#!/usr/bin/env python3
"""Hent ut tekst per slide fra en .pptx som JSON for stil-lint sitt check_slides.

Frittstående: trenger bare python-pptx. Speiler extract_slides i stil-lint
(tittel, brødtekst inkl. tabeller og grupper, notater, skjult-flagg).

    python3 extract_slides.py deck.pptx > slides.json
"""

from __future__ import annotations

import json
import sys


def shape_texts(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from shape_texts(shape.shapes)
            continue
        if getattr(shape, "has_table", False) and shape.has_table:
            rows = []
            for row in shape.table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    rows.append(" | ".join(dict.fromkeys(cells)))
            if rows:
                yield shape, "\n".join(rows)
            continue
        if shape.has_text_frame:
            text = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
            if text:
                yield shape, text


def extract(path: str) -> list[dict]:
    from pptx import Presentation

    out = []
    for i, slide in enumerate(Presentation(path).slides, start=1):
        title, lines = "", []
        title_shape = slide.shapes.title
        for shape, text in shape_texts(slide.shapes):
            if title_shape is not None and shape == title_shape:
                title = text.strip()
            else:
                lines.append(text)
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        out.append({
            "number": i,
            "title": title,
            "body": "\n".join(lines).strip(),
            "notes": notes,
            "hidden": slide._element.get("show") == "0",
        })
    return out


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("bruk: extract_slides.py deck.pptx")
    json.dump(extract(sys.argv[1]), sys.stdout, ensure_ascii=False, indent=1)
    print()
