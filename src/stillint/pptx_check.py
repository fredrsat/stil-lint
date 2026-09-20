"""Sjekk av PowerPoint-presentasjoner.

Hver slide blir ett "avsnitt" i ett dokument (tittel + punkter), slik at
avsnittsregler treffer per slide og dokumentregler (fraktal gjentakelse,
synonymkarusell) ser hele dekket. Speaker notes er prosa og sjekkes separat
med sakprosa-profilen.

Krever python-pptx (pip install stil-lint[pptx]).
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from .engine import Engine


@dataclass
class SlideText:
    number: int                 # 1-basert, som i PowerPoint
    title: str
    body: str
    notes: str
    skipped: bool = False       # AGENT-META o.l.


@dataclass
class DeckReport:
    file: str
    slides: list[SlideText] = field(default_factory=list)
    deck_result: dict | None = None
    notes_result: dict | None = None

    @property
    def verdict(self) -> str:
        verdicts = [r["verdict"] for r in (self.deck_result, self.notes_result) if r]
        if "revise" in verdicts:
            return "revise"
        if "pass_with_notes" in verdicts:
            return "pass_with_notes"
        return "pass"


def extract_slides(path: Path) -> list[SlideText]:
    from pptx import Presentation

    prs = Presentation(str(path))
    slides: list[SlideText] = []
    for i, slide in enumerate(prs.slides, start=1):
        title = ""
        lines: list[str] = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
            if not text:
                continue
            if shape == slide.shapes.title:
                title = text.strip()
            else:
                lines.append(text)
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        body = "\n".join(lines).strip()
        # Skjulte metadata-slides (f.eks. AGENT-META fra agent-meta-pptx) skal ikke lintes.
        skipped = title.upper().startswith("AGENT-META") or body.upper().startswith("AGENT-META")
        slides.append(SlideText(number=i, title=title, body=body, notes=notes, skipped=skipped))
    return slides


def _one_paragraph(text: str) -> str:
    """Klem en slides tekst sammen til ett avsnitt (ingen blanklinjer)."""
    return re.sub(r"\n\s*\n", "\n", text).strip()


async def check_deck(path: Path, mode: str = "fast", engine: Engine | None = None) -> DeckReport:
    engine = engine or Engine()
    slides = extract_slides(path)
    report = DeckReport(file=str(path), slides=slides)

    checkable = [s for s in slides if not s.skipped and (s.title or s.body)]
    if checkable:
        # Tittelen prefikses med "# " slik at overskriftsregler (B07 title case)
        # gjenkjenner den. Tittel og punkter holdes i samme avsnitt så
        # avsnittsindeks == slide.
        deck_text = "\n\n".join(
            _one_paragraph((f"# {s.title}\n" if s.title else "") + s.body)
            for s in checkable
        )
        report.deck_result = await engine.check_text(deck_text, genre="slide", mode=mode)
        # Avsnittsindeks -> slidenummer
        index_to_slide = {idx: s.number for idx, s in enumerate(checkable)}
        for finding in report.deck_result["findings"]:
            if finding.get("paragraph") is not None:
                finding["slide"] = index_to_slide.get(finding["paragraph"])

    notes = [s for s in slides if not s.skipped and s.notes]
    if notes:
        notes_text = "\n\n".join(_one_paragraph(s.notes) for s in notes)
        report.notes_result = await engine.check_text(notes_text, genre="sakprosa", mode=mode)
        index_to_slide = {idx: s.number for idx, s in enumerate(notes)}
        for finding in report.notes_result["findings"]:
            if finding.get("paragraph") is not None:
                finding["slide"] = index_to_slide.get(finding["paragraph"])

    return report


def format_report(report: DeckReport) -> str:
    lines = [f"{report.verdict.upper()}  {report.file}"]
    for label, result in (("slides", report.deck_result), ("notater", report.notes_result)):
        if not result:
            continue
        lines.append(f"  [{label}] {result['verdict']} score={result['score']}")
        for f in result["findings"]:
            where = f"slide {f['slide']}" if f.get("slide") is not None else "hele dekket"
            note = " (råd)" if f.get("advisory") else ""
            lines.append(f"    [{f['rule']}] {where}, p={f['p']}, sev={f['severity']}{note}")
            lines.append(f"        {f['hint']}")
            if f.get("evidence"):
                lines.append(f"        treff: {f['evidence']}")
        for m in result["missing"]:
            lines.append(f"    MANGLER: {m}")
    return "\n".join(lines)
