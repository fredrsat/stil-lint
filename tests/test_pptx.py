import asyncio

import pytest

pptx_lib = pytest.importorskip("pptx")

from stillint.engine import Engine  # noqa: E402
from stillint.pptx_check import check_deck, extract_slides  # noqa: E402


def build_deck(path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    layout = prs.slide_layouts[1]  # tittel + innhold

    s1 = prs.slides.add_slide(layout)
    s1.shapes.title.text = "Kvartalsrapport Q3"
    s1.placeholders[1].text = ("Salget økte 12 % til 4,2 mill. kr\n"
                               "Tre nye kunder i Bergen\nLevering i snitt 2,1 dager")
    s1.notes_slide.notes_text_frame.text = (
        "Si at hovedtallet er 12 % vekst, drevet av avtalen med Vestkraft i august.")

    s2 = prs.slides.add_slide(layout)
    s2.shapes.title.text = "En Sømløs Og Banebrytende Reise"
    s2.placeholders[1].text = ("Vår sømløse og banebrytende plattform\n"
                               "Skreddersydd og helhetlig verdiskaping\n"
                               "Nøkkelen til synergi")
    s2.notes_slide.notes_text_frame.text = (
        "Det er viktig å merke seg at dette er nøkkelen til suksess. Håper dette hjelper!")

    meta = prs.slides.add_slide(prs.slide_layouts[5])
    meta.shapes.title.text = "AGENT-META"
    tb = meta.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(2))
    tb.text_frame.text = "sømløs banebrytende metadata som ikke skal lintes"

    prs.save(str(path))
    return path


def test_extract_slides(tmp_path):
    deck = build_deck(tmp_path / "deck.pptx")
    slides = extract_slides(deck)
    assert len(slides) == 3
    assert slides[0].title == "Kvartalsrapport Q3"
    assert "12 %" in slides[0].body
    assert "Vestkraft" in slides[0].notes
    assert slides[2].skipped  # AGENT-META hoppes over


def test_check_deck_flags_slop_slide(tmp_path):
    deck = build_deck(tmp_path / "deck.pptx")
    report = asyncio.run(check_deck(deck, mode="fast", engine=Engine(bank_path=tmp_path / "b.db")))
    assert report.verdict == "revise"
    deck_rules = {(f["rule"], f.get("slide")) for f in report.deck_result["findings"]}
    assert ("A02_stilord_nb", 2) in deck_rules          # stilord på slide 2
    assert ("B07_title_case_no", None) in deck_rules or any(
        r == "B07_title_case_no" for r, _ in deck_rules)
    assert all(s != 1 for r, s in deck_rules if s)      # den rene sliden flagges ikke
    notes_rules = {f["rule"] for f in report.notes_result["findings"]}
    assert "A04_chatbotrester" in notes_rules            # "Håper dette hjelper!" i notatene


def test_clean_deck_passes(tmp_path):
    from pptx import Presentation

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Salget økte 12 % i tredje kvartal"
    s.placeholders[1].text = "4,2 mill. kr omsetning\nTre nye kunder i Bergen"
    path = tmp_path / "clean.pptx"
    prs.save(str(path))

    report = asyncio.run(check_deck(path, mode="fast", engine=Engine(bank_path=tmp_path / "b.db")))
    assert report.deck_result["findings"] == []
